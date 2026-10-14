import io
import zipfile
from uuid import UUID

import pytest
from sqlalchemy import select

from app.core.database import get_session_factory
from app.core.unit_of_work import UnitOfWork
from app.infrastructure.documents import PerenderBerkas
from app.modules.laporan.application.ekspor import LayananEkspor
from app.modules.laporan.infrastructure.models import EksporDokumen
from app.modules.laporan.infrastructure.penyusun_cuplikan import PenyusunCuplikanDb
from tests.fabrik import buka_siklus, instrumen_infokom, instrumen_lamsama, tahun_akademik

pytestmark = pytest.mark.db


async def bukti_dan_dokumen(dunia, ta, prodi_ids, responden=200, populasi=300):
    bk = await dunia.post(
        "sf", "/bukti", json={"jenis": "tautan", "judul": "Laporan tracer", "url": "https://dir.kampus.ac.id/t", "akses_tautan": "publik"}
    )
    r = await dunia.post(
        "sf",
        "/dokumen-eksternal",
        json={
            "jenis": "tracer_study",
            "judul": "Tracer 2025",
            "tahun_akademik_id": str(ta),
            "populasi": populasi,
            "jumlah_responden": responden,
            "bukti_id": bk.json()["id"],
            "prodi_id": prodi_ids,
        },
    )
    assert r.status_code == 201, r.text
    return r.json()


# ------------------------------------------------------------------ dokumen eksternal
async def test_dokumen_tracer_dikelola_staf_fakultas_dibaca_prodi_us_e1(dunia, prodi_id):
    ta = (await tahun_akademik())[-1]
    kimia = str(await prodi_id("KIMIA"))
    d = await bukti_dan_dokumen(dunia, ta, [kimia])
    assert d["butir_dibebaskan"] == 0 and d["bukti"]["judul"] == "Laporan tracer" and d["prodi"][0]["kode"] == "KIMIA"
    for who in ("k_kimia", "sp_kimia", "dekan"):
        assert (await dunia.post(who, "/dokumen-eksternal", json={})).status_code == 403, who
    assert (await dunia.get("k_kimia", "/dokumen-eksternal")).json()["meta"]["total"] == 1
    assert (await dunia.get("k_ilkom", "/dokumen-eksternal")).json()["meta"]["total"] == 0  # dokumen hanya untuk prodi sasaran
    assert (await dunia.get("sf", f"/dokumen-eksternal?prodi_id={kimia}")).json()["meta"]["total"] == 1
    # responden tidak boleh melebihi populasi
    bk = (await dunia.post("sf", "/bukti", json={"jenis": "tautan", "judul": "Bukti B", "url": "https://x.kampus.ac.id/b"})).json()["id"]
    r = await dunia.post(
        "sf",
        "/dokumen-eksternal",
        json={
            "jenis": "tracer_study",
            "judul": "Salah",
            "tahun_akademik_id": str(ta),
            "populasi": 10,
            "jumlah_responden": 20,
            "bukti_id": bk,
            "prodi_id": [kimia],
        },
    )
    assert r.status_code == 422


async def test_indikator_responden_kecil_tampil_tidak_memadai_ec16(dunia, prodi_id):
    ta = (await tahun_akademik())[-1]
    kimia = str(await prodi_id("KIMIA"))
    d = await bukti_dan_dokumen(dunia, ta, [kimia], responden=12, populasi=400)
    r = await dunia.put(
        "sf",
        f"/dokumen-eksternal/{d['id']}/indikator",
        json=[
            {"prodi_id": kimia, "kode": "masa_tunggu_bulan", "nama": "Masa tunggu", "nilai": 3.1, "satuan": "bulan"},
            {"prodi_id": kimia, "kode": "kepuasan", "nama": "Kepuasan", "nilai": 3.9, "satuan": "skala 4", "jumlah_responden": 150},
        ],
    )
    assert r.status_code == 200
    ind = {i["kode"]: i for i in (await dunia.get("k_kimia", f"/prodi/{kimia}/indikator-eksternal")).json()}
    assert (
        ind["masa_tunggu_bulan"]["data_memadai"] is False and ind["masa_tunggu_bulan"]["nilai"] is None
    )  # angka disembunyikan, bukan menyesatkan
    assert ind["kepuasan"]["data_memadai"] is True and float(ind["kepuasan"]["nilai"]) == 3.9
    assert (await dunia.get("k_ilkom", f"/prodi/{kimia}/indikator-eksternal")).status_code == 403


async def test_impor_indikator_melaporkan_baris_gagal_ec17(dunia, prodi_id):
    ta = (await tahun_akademik())[-1]
    kimia = str(await prodi_id("KIMIA"))
    d = await bukti_dan_dokumen(dunia, ta, [kimia])
    csv = (
        b"prodi_kode,kode,nama,nilai,satuan,jumlah_responden\n"
        b"KIMIA,masa_tunggu_bulan,Masa tunggu,4.2,bulan,200\n"
        b"KIMIA,kesesuaian,Kesesuaian bidang,bukan-angka,persen,200\n"
        b"XXXX,kepuasan,Kepuasan,3.5,skala 4,200\n"
        b"KIMIA,kepuasan,Kepuasan,3.5,skala 4,200\n"
    )
    r = await dunia.post("sf", f"/dokumen-eksternal/{d['id']}/indikator/impor", files={"berkas": ("i.csv", csv, "text/csv")})
    assert r.status_code == 202 and r.json()["baris_berhasil"] == 2
    assert {g["baris"] for g in r.json()["baris_gagal"]} == {3, 4}
    assert len((await dunia.get("sf", f"/prodi/{kimia}/indikator-eksternal")).json()) == 2


async def test_butir_menunggu_dokumen_direktorat_dibebaskan_saat_indikator_tiba_ec15(dunia, prodi_id):
    ta = await tahun_akademik()
    await instrumen_lamsama(dekan_id=dunia.id["dekan"])
    sk = await buka_siklus(dunia, "KIMIA", ta[-1])
    kimia = str(await prodi_id("KIMIA"))

    async def menunggu() -> int:
        return (await dunia.get("dekan", f"/siklus/{sk['id']}/butir?menunggu_dokumen_eksternal=true")).json()["meta"]["total"]

    assert await menunggu() == 1
    d = await bukti_dan_dokumen(dunia, ta[-1], [kimia])
    assert await menunggu() == 1  # dokumen saja belum cukup; indikator belum diisi
    r = await dunia.put(
        "sf",
        f"/dokumen-eksternal/{d['id']}/indikator",
        json=[{"prodi_id": kimia, "kode": "masa_tunggu_bulan", "nama": "Masa tunggu", "nilai": 4, "satuan": "bulan"}],
    )
    assert r.json()["butir_dibebaskan"] == 1 and await menunggu() == 0


# ------------------------------------------------------------------ penilaian
async def isi_skor(dunia, who, sid, skor: dict[str, float]):
    for b in (await dunia.get(who, f"/siklus/{sid}/butir")).json()["data"]:
        d = (await dunia.get(who, f"/butir-siklus/{b['id']}")).json()
        r = await dunia.patch(
            who,
            f"/butir-siklus/{b['id']}",
            json={"row_version": d["row_version"], "skor_estimasi": skor[b["butir"]["kode"]], "status_kelengkapan": "lengkap"},
        )
        assert r.status_code == 200, r.text


async def test_simulasi_lam_infokom_nilai_berbobot_dan_label_estimasi_us_f4(dunia):
    ta = await tahun_akademik()
    await instrumen_infokom(dekan_id=dunia.id["dekan"])
    sk = await buka_siklus(dunia, "ILKOM", ta[-1], target_peringkat_kode="unggul_3")
    # butir: A.1(1) 1.1(5) 1.2(5) 2.1(10) 3.1(9) → Σ bobot 30; semua skor 4 → 120; semua 3.3 → 99
    await isi_skor(dunia, "k_ilkom", sk["id"], {"A.1": 4, "1.1": 4, "1.2": 4, "2.1": 4, "3.1": 4})
    r = await dunia.post("k_ilkom", f"/siklus/{sk['id']}/simulasi-skor")
    assert r.status_code == 201
    j = r.json()
    assert j["estimasi_internal"] is True and "bukan penilaian resmi" in j["peringatan"] and j["dihitung_pada"]
    assert float(j["nilai_akhir"]) == 120.0 and j["nilai_akhir_maks"] == "400.00" or float(j["nilai_akhir_maks"]) == 400.0
    assert j["prediksi_peringkat"]["kode"] == "tidak_terakreditasi"  # 120 < 200 pada skala 400 (instrumen uji hanya memuat Σ bobot 30)
    assert j["target_peringkat"]["kode"] == "unggul_3" and j["di_bawah_target"] is True
    kriteria = {k["kode"]: k for k in j["rincian"]["kriteria"]}
    assert kriteria["1"]["rerata"] == 4.0
    assert j["rincian"]["prioritas_perbaikan"] == [] or isinstance(j["rincian"]["prioritas_perbaikan"], list)
    # GET mengembalikan hasil terakhir; Staf tidak berhak
    assert (await dunia.get("k_ilkom", f"/siklus/{sk['id']}/simulasi-skor")).json()["nilai_akhir"] == j["nilai_akhir"]
    assert (await dunia.get("sp_ilkom", f"/siklus/{sk['id']}/simulasi-skor")).status_code == 403
    assert (await dunia.post("k_kimia", f"/siklus/{sk['id']}/simulasi-skor")).status_code == 403


async def test_simulasi_lamsama_ambang_dan_syarat_perlu_us_f4(dunia):
    ta = await tahun_akademik()
    await instrumen_lamsama(dekan_id=dunia.id["dekan"])
    sk = await buka_siklus(dunia, "KIMIA", ta[-1], target_peringkat_kode="unggul_2")
    await isi_skor(dunia, "k_kimia", sk["id"], {"1": 2, "2": 4, "3": 4})  # indikator 1 = syarat perlu hanya 2 (<3 untuk unggul_2)
    j = (await dunia.post("k_kimia", f"/siklus/{sk['id']}/simulasi-skor")).json()
    assert j["nilai_akhir"] is None and j["prediksi_peringkat"]["kode"] == "terakreditasi"
    u2 = next(p for p in j["rincian"]["per_peringkat"] if p["kode"] == "unggul_2")
    assert (
        u2["tercapai"] is False
        and u2["butir_di_bawah_ambang"][0]["kode"] == "1"
        and u2["butir_di_bawah_ambang"][0]["is_syarat_perlu"] is True
    )
    assert j["rincian"]["prioritas_perbaikan"][0]["kode"] == "1"  # kriteria di bawah target masuk daftar prioritas


async def test_dasbor_fakultas_progres_risiko_dan_filter_us_f1(dunia, prodi_id):
    ta = await tahun_akademik()
    await instrumen_lamsama(dekan_id=dunia.id["dekan"])
    await instrumen_infokom(dekan_id=dunia.id["dekan"])
    kimia = await buka_siklus(dunia, "KIMIA", ta[-1])
    await buka_siklus(dunia, "ILKOM", ta[-1])
    await isi_skor(dunia, "k_kimia", kimia["id"], {"1": 3, "2": 3, "3": 3})
    d = (await dunia.get("dekan", "/dasbor/fakultas")).json()
    peta = {p["prodi"]["kode"]: p for p in d["prodi"]}
    assert set(peta) == {"ILKOM", "KIMIA", "AKTUARIA"}
    assert peta["KIMIA"]["persen_siap"] == 100.0 and peta["KIMIA"]["butir"]["lengkap"] == 3
    assert peta["ILKOM"]["persen_siap"] == 0.0 and peta["ILKOM"]["tingkat_risiko"] in ("kuning", "merah")
    assert peta["AKTUARIA"]["siklus_id"] is None and peta["AKTUARIA"]["tingkat_risiko"] == "abu"
    assert peta["KIMIA"]["simulasi"]["estimasi_internal"] is True
    assert (await dunia.get("sf", "/dasbor/fakultas")).status_code == 200
    for who in ("k_kimia", "sp_kimia", "it"):
        assert (await dunia.get(who, "/dasbor/fakultas")).status_code == 403, who
    # dasbor prodi: Kaprodi hanya prodinya; daftar butir belum lengkap tersedia (monitoring berjalan)
    dp = (await dunia.get("k_ilkom", f"/dasbor/prodi/{await prodi_id('ILKOM')}")).json()
    assert len(dp["butir_belum_lengkap"]) == 5
    assert (await dunia.get("k_ilkom", f"/dasbor/prodi/{await prodi_id('KIMIA')}")).status_code == 403
    one = (await dunia.get("dekan", f"/dasbor/fakultas?prodi_id={await prodi_id('KIMIA')}")).json()["prodi"]
    assert len(one) == 1


async def test_gap_analysis_dan_benchmark_hanya_dekan_dan_indikator_setara_ec20(dunia, prodi_id):
    ta = await tahun_akademik()
    kimia, ilkom = await prodi_id("KIMIA"), await prodi_id("ILKOM")
    # dua prodi, instrumen berbeda, data inti setara → perbandingan memakai indikator inti, bukan skor mentah
    for kode, pid, n_pub in (("k_kimia", kimia, 3), ("k_ilkom", ilkom, 1)):
        d = await dunia.post(
            kode,
            "/dosen",
            json={
                "nama": f"Dosen {kode}",
                "pendidikan_tertinggi": "s3",
                "status_kepegawaian": "tetap",
                "afiliasi": [{"prodi_id": str(pid), "tahun_akademik_id": str(ta[-1])}],
            },
        )
        assert d.status_code == 201
        for i in range(n_pub):
            r = await dunia.post(
                kode,
                "/publikasi?abaikan_kemiripan=true",
                json={
                    "judul": f"Publikasi {kode} nomor {i} unik",
                    "jenis": "jurnal",
                    "tahun_terbit": 2025,
                    "tahun_akademik_id": str(ta[-1]),
                },
            )
            assert r.status_code == 201, r.text
    r = await dunia.get("dekan", f"/analisis/benchmark?indikator=publikasi_per_dosen&tahun_akademik_id={ta[-1]}")
    assert r.status_code == 200
    data = {x["prodi"]: float(x["nilai"]) if x["nilai"] is not None else None for x in r.json()["data"]}
    assert data["KIMIA"] == 3.0 and data["ILKOM"] == 1.0 and data["AKTUARIA"] is None
    assert r.json()["data"][0]["prodi"] == "KIMIA"  # urut menurut makin_tinggi_makin_baik
    assert r.json()["indikator"]["makin_tinggi_makin_baik"] is True
    for who in ("k_kimia", "sf", "it"):
        assert (await dunia.get(who, f"/analisis/benchmark?indikator=publikasi_per_dosen&tahun_akademik_id={ta[-1]}")).status_code == 403
    assert (await dunia.get("dekan", f"/analisis/benchmark?indikator=skor_mentah_lintas_lam&tahun_akademik_id={ta[-1]}")).status_code == 404
    # gap: tepat salah satu parameter
    assert (await dunia.get("dekan", "/analisis/gap")).status_code == 422
    await instrumen_lamsama(dekan_id=dunia.id["dekan"])
    sk = await buka_siklus(dunia, "KIMIA", ta[-1], target_peringkat_kode="unggul_2")
    await isi_skor(dunia, "k_kimia", sk["id"], {"1": 3, "2": 2, "3": 3})
    gap = (await dunia.get("dekan", f"/analisis/gap?siklus_id={sk['id']}")).json()
    assert gap["butir_di_bawah_target"] == 1 and gap["butir"][0]["kode"] == "2" and float(gap["butir"][0]["selisih"]) == -1.0


# ------------------------------------------------------------------ ekspor (renderer nyata, penyimpanan palsu)
async def test_proses_ekspor_docx_xlsx_draf_dan_final(dunia, storage, antrean):
    ta = await tahun_akademik()
    await instrumen_lamsama(dekan_id=dunia.id["dekan"])
    sk = await buka_siklus(dunia, "KIMIA", ta[-1])
    await isi_skor(dunia, "k_kimia", sk["id"], {"1": 3, "2": 3, "3": 3})
    butir_pertama = (await dunia.get("k_kimia", f"/siklus/{sk['id']}/butir")).json()["data"][0]["id"]
    (await dunia.put("k_kimia", f"/butir-siklus/{butir_pertama}/narasi/utama", json={"isi": "Narasi untuk dokumen."})).raise_for_status()

    async def jalankan(jenis: str, fmt: str, final: bool = False) -> EksporDokumen:
        r = await dunia.post("k_kimia", f"/siklus/{sk['id']}/ekspor", json={"jenis": jenis, "format": fmt, "final": final})
        assert r.status_code == 202, r.text
        async with get_session_factory()() as s:
            await LayananEkspor(UnitOfWork(s), PenyusunCuplikanDb(s), storage, PerenderBerkas()).proses(UUID(r.json()["id"]))
        st = (await dunia.get("k_kimia", f"/ekspor/{r.json()['id']}")).json()
        assert st["status"] == "selesai", st
        return st

    docx = await jalankan("led", "docx")
    assert docx["metadata"]["final"] is False and docx["ukuran_byte"] > 5000
    key = next(k for k in storage.objek if k.endswith("-led-draf.docx"))
    z = zipfile.ZipFile(io.BytesIO(storage.objek[key]))
    teks = z.read("word/document.xml").decode() + z.read("word/header1.xml").decode()
    assert "Narasi untuk dokumen." in teks and "DRAF" in teks and "Indikator 1 uji" in teks and "estimasi internal" in teks
    xlsx = await jalankan("data_kuantitatif", "xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(storage.objek[next(k for k in storage.objek if k.endswith(".xlsx"))]))
    assert set(wb.sheetnames) == {"Ringkasan", "Butir", "Data terpetakan", "Bukti"} and wb["Butir"].max_row == 4
    assert xlsx["metadata"]["prodi"] == "Kimia"
    # unduh: URL sementara dari penyimpanan privat
    u = await dunia.get("k_kimia", f"/ekspor/{xlsx['id']}/unduh")
    assert u.status_code == 200 and u.json()["url"].startswith("http://storage.test/unduh/ekspor/")
    # gap_analysis hanya Dekan
    assert (await dunia.post("k_kimia", f"/siklus/{sk['id']}/ekspor", json={"jenis": "gap_analysis", "format": "xlsx"})).status_code == 403
    assert (await dunia.post("dekan", f"/siklus/{sk['id']}/ekspor", json={"jenis": "gap_analysis", "format": "xlsx"})).status_code == 202


def test_html_led_per_lam_berbeda_namun_data_sama():
    """Templat tiap LAM dirender terpisah (US-D3) dari cuplikan yang sama."""
    from app.infrastructure.documents.pdf_renderer import render_html

    def cup(lam: str) -> dict:
        return {
            "instrumen": {"lam": lam, "kode_versi": "x", "skor_maks": 4, "nilai_akhir_maks": None},
            "prodi": {"nama": "Kimia", "jenjang": "S1"},
            "siklus": {"nama": "S", "ts": "2025/2026"},
            "skor": {},
            "bukti_ids": [],
            "kriteria": [
                {
                    "kode": "1",
                    "nama": "Budaya Mutu",
                    "bobot": 40,
                    "butir": [
                        {
                            "kode": "1.1",
                            "elemen": "Elemen",
                            "skor_estimasi": 3,
                            "bobot": 2,
                            "jenis_ipo": "input",
                            "tahap_ppepp": "penetapan",
                            "status_kelengkapan": "lengkap",
                            "narasi": [],
                            "skema_narasi": [],
                            "ambang": [],
                            "bukti": [{"judul": "B", "url": "https://x", "status_cek": "tidak_terjangkau"}],
                            "data_terpetakan": [{"sumber": "publikasi", "jumlah": 2, "kosong": False, "nilai": None}],
                        }
                    ],
                }
            ],
        }

    meta = {"data_per": "2026-10-04", "pengajuan_nomor": 2}
    infokom = render_html(cup("LAM_INFOKOM"), final=True, metadata=meta)
    lamsama = render_html(cup("LAMSAMA"), final=False, metadata=meta)
    assert "Bobot" in infokom and "(bobot 40)" in infokom and "watermark" not in infokom
    assert "Syarat perlu" in lamsama and "Indikator 1.1" in lamsama and "watermark" in lamsama
    assert "TIDAK TERJANGKAU" in infokom and "publikasi: 2 catatan" in infokom


# ------------------------------------------------------------------ LLM
class KlienPalsu:
    def __init__(self, gagal: bool = False):
        self.gagal, self.panggilan = gagal, []

    async def uji(self, base_url, model, api_key):
        if self.gagal:
            raise RuntimeError("koneksi ditolak")
        return "OK (siap)"

    async def selesaikan(self, base_url, model, api_key, pesan, parameter):
        from app.modules.llm.domain.ports import HasilLlm

        self.panggilan.append((base_url, api_key, pesan))
        if self.gagal:
            raise RuntimeError("timeout")
        return HasilLlm("Program studi telah menetapkan standar mutu penelitian ...", 120, 80)


async def test_llm_offline_default_online_butuh_persetujuan_dan_draf_ditinjau_manusia(dunia, monkeypatch, antrean):
    from app.modules.llm.application.layanan import LayananLlm

    klien = KlienPalsu()
    monkeypatch.setattr("app.modules.llm.api.router.KlienOpenAICompatible", lambda: klien)
    ta = await tahun_akademik()
    await instrumen_lamsama(dekan_id=dunia.id["dekan"])
    sk = await buka_siklus(dunia, "KIMIA", ta[-1])
    butir = (await dunia.get("k_kimia", f"/siklus/{sk['id']}/butir")).json()["data"][0]["id"]
    # tanpa konfigurasi aktif → 503, penulisan manual tetap berjalan (EC-24)
    r = await dunia.post("k_kimia", f"/butir-siklus/{butir}/narasi/utama/draf-llm", json={})
    assert r.status_code == 503 and r.json()["error"]["code"] == "LLM_TIDAK_TERSEDIA"
    assert (await dunia.put("k_kimia", f"/butir-siklus/{butir}/narasi/utama", json={"isi": "Manual tetap bisa."})).status_code == 200
    # konfigurasi: kunci API tidak pernah dikembalikan; hanya Staf Fakultas menyiapkan, Dekan mengaktifkan
    body = {"nama": "Model lokal", "mode": "offline", "base_url": "http://llm:11434/v1", "model": "lokal-1", "api_key": "rahasia-sekali"}
    assert (await dunia.post("k_kimia", "/konfigurasi-llm", json=body)).status_code == 403
    k = (await dunia.post("sf", "/konfigurasi-llm", json=body)).json()
    assert k["punya_api_key"] is True and "api_key" not in k and "rahasia-sekali" not in str(k)
    async with get_session_factory()() as s:
        from app.modules.llm.infrastructure.models import KonfigurasiLlm

        mentah = (await s.execute(select(KonfigurasiLlm.api_key_terenkripsi))).scalar_one()
        assert b"rahasia-sekali" not in mentah  # terenkripsi di basis data
    assert (await dunia.post("dekan", f"/konfigurasi-llm/{k['id']}/aktifkan", json={})).status_code == 409  # wajib uji koneksi dulu
    assert (await dunia.post("sf", f"/konfigurasi-llm/{k['id']}/uji")).json()["berhasil"] is True
    assert (await dunia.post("sf", f"/konfigurasi-llm/{k['id']}/aktifkan", json={})).status_code == 403
    assert (await dunia.post("dekan", f"/konfigurasi-llm/{k['id']}/aktifkan", json={})).json()["is_aktif"] is True

    # mode online: butuh konfirmasi eksplisit Dekan
    online = (
        await dunia.post(
            "sf", "/konfigurasi-llm", json={**body, "nama": "Penyedia daring", "mode": "online", "base_url": "https://api.penyedia.com/v1"}
        )
    ).json()
    (await dunia.post("sf", f"/konfigurasi-llm/{online['id']}/uji")).raise_for_status()
    r = await dunia.post("dekan", f"/konfigurasi-llm/{online['id']}/aktifkan", json={})
    assert r.status_code == 422 and r.json()["error"]["details"][0]["kolom"] == "konfirmasi_online"
    assert (await dunia.post("dekan", f"/konfigurasi-llm/{k['id']}/nonaktifkan")).status_code == 200

    # permintaan draf (offline aktif lagi) → antrean 'llm'; hasil diproses worker
    (await dunia.post("dekan", f"/konfigurasi-llm/{k['id']}/aktifkan", json={})).raise_for_status()
    r = await dunia.post("k_kimia", f"/butir-siklus/{butir}/narasi/utama/draf-llm", json={"arahan": "Tekankan tiga tahun terakhir."})
    assert r.status_code == 202 and r.json()["mode"] == "offline" and r.json()["model"] == "lokal-1"
    assert antrean[-1][0] == "simak.draf_llm"
    assert (await dunia.post("sp_kimia", f"/butir-siklus/{butir}/narasi/utama/draf-llm", json={})).status_code == 403  # hanya Kaprodi
    rid = r.json()["id"]
    async with get_session_factory()() as s:
        await LayananLlm(UnitOfWork(s), klien).proses(UUID(rid), "Tekankan tiga tahun terakhir.", True)
    hasil = (await dunia.get("k_kimia", f"/permintaan-llm/{rid}")).json()
    assert hasil["status"] == "selesai" and hasil["hasil"].startswith("Program studi") and "wajib ditinjau" in hasil["pengingat"]
    assert klien.panggilan[0][2][0]["role"] == "system" and "Tekankan tiga tahun terakhir." in klien.panggilan[0][2][1]["content"]
    # draf tidak otomatis masuk narasi; disimpan Kaprodi → sumber llm / llm_disunting
    narasi = (await dunia.get("k_kimia", f"/butir-siklus/{butir}")).json()["narasi"][0]
    assert narasi["isi"] == "Manual tetap bisa." and narasi["sumber"] == "manual"
    r = await dunia.put(
        "k_kimia",
        f"/butir-siklus/{butir}/narasi/utama",
        json={"row_version": narasi["row_version"], "isi": hasil["hasil"], "permintaan_llm_id": rid},
    )
    assert r.json()["sumber"] == "llm"
    r = await dunia.put(
        "k_kimia",
        f"/butir-siklus/{butir}/narasi/utama",
        json={"row_version": 2, "isi": hasil["hasil"] + " Ditambah Kaprodi.", "permintaan_llm_id": rid},
    )
    assert r.json()["sumber"] == "llm_disunting"
    # permintaan milik butir lain tidak dapat dipakai
    butir_lain = (await dunia.get("k_kimia", f"/siklus/{sk['id']}/butir")).json()["data"][1]["id"]
    assert (
        await dunia.put("k_kimia", f"/butir-siklus/{butir_lain}/narasi/utama", json={"isi": "x", "permintaan_llm_id": rid})
    ).status_code == 422


async def test_llm_gagal_tidak_menghilangkan_data_ec24(dunia, monkeypatch, antrean):
    from app.modules.llm.application.layanan import LayananLlm

    klien = KlienPalsu(gagal=True)
    monkeypatch.setattr("app.modules.llm.api.router.KlienOpenAICompatible", lambda: klien)
    ta = await tahun_akademik()
    await instrumen_lamsama(dekan_id=dunia.id["dekan"])
    sk = await buka_siklus(dunia, "KIMIA", ta[-1])
    butir = (await dunia.get("k_kimia", f"/siklus/{sk['id']}/butir")).json()["data"][0]["id"]
    async with get_session_factory()() as s:
        from app.modules.llm.domain.ports import ModeLlm
        from app.modules.llm.infrastructure.models import KonfigurasiLlm

        s.add(KonfigurasiLlm(nama="Lokal", mode=ModeLlm.OFFLINE, base_url="http://llm:11434/v1", model="m", is_aktif=True, diuji_pada=None))
        await s.commit()
    (
        await dunia.put("k_kimia", f"/butir-siklus/{butir}/narasi/utama", json={"isi": "Tulisan manual yang harus selamat."})
    ).raise_for_status()
    rid = (await dunia.post("k_kimia", f"/butir-siklus/{butir}/narasi/utama/draf-llm", json={})).json()["id"]
    async with get_session_factory()() as s:
        await LayananLlm(UnitOfWork(s), klien).proses(UUID(rid), None, True)
    st = (await dunia.get("k_kimia", f"/permintaan-llm/{rid}")).json()
    assert st["status"] == "gagal" and "timeout" in st["galat"] and st["hasil"] is None
    assert (await dunia.get("k_kimia", f"/butir-siklus/{butir}")).json()["narasi"][0]["isi"] == "Tulisan manual yang harus selamat."
