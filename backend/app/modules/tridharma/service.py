"""Layanan CRUD generik data inti: cakupan prodi, penguncian laporan, versi optimistik, duplikasi, bukti."""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError
from sqlalchemy import Select, delete, func, inspect, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm.exc import StaleDataError

from app.core import audit
from app.core.exceptions import (
    AksesDitolak,
    Duplikat,
    GalatDomain,
    KonflikStatus,
    KonflikVersi,
    TidakDitemukan,
    ValidasiGagal,
)
from app.core.pagination import ParamHalaman, halaman
from app.core.permissions import SF, K, Pemanggil
from app.core.unit_of_work import UnitOfWork
from app.infrastructure.documents.templat_impor import TIPE_CSV, TIPE_XLSX, buat_csv, buat_xlsx
from app.modules.bukti import public as bukti
from app.modules.instrumen import public as instrumen
from app.modules.instrumen.public import JenisEntitas
from app.modules.laporan import public as laporan
from app.modules.referensi import public as referensi
from app.modules.tridharma import models as m
from app.modules.tridharma.registry import Spec
from app.modules.tridharma.templat import kolom_impor

KOLOM_INTERNAL = {"deleted_at"}


def _sekarang() -> datetime:
    return datetime.now(UTC)


class LayananDataInti:
    def __init__(self, spec: Spec, uow: UnitOfWork):
        self.spec, self.uow, self.session = spec, uow, uow.session
        self.model = spec.model
        self.kolom_tabel = {c.key for c in inspect(spec.model).mapper.column_attrs}
        self.bidang_ekstra = set(spec.skema.model_fields) - self.kolom_tabel

    # ------------------------------------------------------------ util
    def _dasar(self) -> Select[Any]:
        stmt = select(self.model)
        if "deleted_at" in self.kolom_tabel:
            stmt = stmt.where(self.model.deleted_at.is_(None))
        for k, v in self.spec.tetap.items():
            stmt = stmt.where(getattr(self.model, k) == v)
        return stmt

    async def _ke_dict(self, obj: Any, prodi_peta: dict[UUID, Any] | None = None) -> dict[str, Any]:
        d = {k: getattr(obj, k) for k in self.kolom_tabel if k not in KOLOM_INTERNAL}
        prodi_ids = await self.spec.penangan.prodi_ids(self.session, obj)
        peta = prodi_peta or {p.id: p for p in await referensi.daftar_prodi(self.session)}
        d["prodi"] = [{"id": p, "kode": peta[p].kode, "nama": peta[p].nama} for p in prodi_ids if p in peta]
        d.update(await self.spec.penangan.ekstra(self.session, obj))
        d.update(await self._ekstra_dosen(obj))
        return d

    async def _ekstra_dosen(self, obj: Any) -> dict[str, Any]:
        if isinstance(obj, m.KegiatanTridharma):
            baris = (await self.session.execute(select(m.KegiatanDosen).where(m.KegiatanDosen.kegiatan_id == obj.id))).scalars().all()
            return {"dosen": [{"dosen_id": b.dosen_id, "peran": b.peran} for b in baris]}
        if isinstance(obj, m.Publikasi):
            baris = (
                (
                    await self.session.execute(
                        select(m.PublikasiDosen).where(m.PublikasiDosen.publikasi_id == obj.id).order_by(m.PublikasiDosen.urutan_penulis)
                    )
                )
                .scalars()
                .all()
            )
            return {
                "penulis": [
                    {"dosen_id": b.dosen_id, "urutan_penulis": b.urutan_penulis, "is_korespondensi": b.is_korespondensi} for b in baris
                ]
            }
        return {}

    async def _simpan_dosen_terkait(self, obj: Any, data: dict[str, Any]) -> None:
        if isinstance(obj, m.KegiatanTridharma) and "dosen" in data:
            await self.session.execute(delete(m.KegiatanDosen).where(m.KegiatanDosen.kegiatan_id == obj.id))
            for d in data["dosen"]:
                self.session.add(m.KegiatanDosen(kegiatan_id=obj.id, dosen_id=d.dosen_id, peran=d.peran))
        if isinstance(obj, m.Publikasi) and "penulis" in data:
            await self.session.execute(delete(m.PublikasiDosen).where(m.PublikasiDosen.publikasi_id == obj.id))
            for p in data["penulis"]:
                self.session.add(
                    m.PublikasiDosen(
                        publikasi_id=obj.id, dosen_id=p.dosen_id, urutan_penulis=p.urutan_penulis, is_korespondensi=p.is_korespondensi
                    )
                )

    # ------------------------------------------------------------ otorisasi
    async def _muat(self, obj_id: UUID) -> Any:
        obj = (await self.session.execute(self._dasar().where(self.model.id == obj_id))).scalar_one_or_none()
        if obj is None:
            raise TidakDitemukan()
        return obj

    async def _wajib_baca(self, pemanggil: Pemanggil, obj: Any) -> set[UUID]:
        prodi = await self.spec.penangan.prodi_ids(self.session, obj)
        cakupan = pemanggil.cakupan_baca()
        if cakupan.semua:
            return prodi
        # Data tingkat UPPS (Langsung boleh_kosong) terbaca oleh Kaprodi/Staf Prodi; data bersama tanpa prodi tidak.
        if not prodi and getattr(self.spec.penangan, "boleh_kosong", False):
            return prodi
        if not any(cakupan.mengizinkan(p) for p in prodi):
            raise AksesDitolak(alasan="di_luar_cakupan_prodi")
        return prodi

    def _boleh_tulis(self, pemanggil: Pemanggil, prodi: set[UUID], obj: Any | None) -> bool:
        if self.spec.kolom_hanya_sf and obj is not None and getattr(obj, self.spec.kolom_hanya_sf) and not pemanggil.punya(SF):
            return False
        if pemanggil.punya(SF) and self.spec.sf_bebas:
            return True
        if not prodi:
            return pemanggil.punya(SF)
        cakupan = pemanggil.cakupan_tulis()
        return any(cakupan.mengizinkan(p) for p in prodi)

    async def _wajib_tulis(self, pemanggil: Pemanggil, prodi: set[UUID], obj: Any | None = None) -> None:
        if not self._boleh_tulis(pemanggil, prodi, obj):
            raise AksesDitolak(alasan="di_luar_cakupan_prodi")
        if len(prodi) == 1:  # data yang hanya teratribusi ke satu prodi ikut terkunci (EC-07, EC-09)
            await laporan.wajib_tidak_terkunci(self.session, next(iter(prodi)))

    # ------------------------------------------------------------ baca
    async def daftar(
        self,
        pemanggil: Pemanggil,
        param: ParamHalaman,
        *,
        prodi_id: UUID | None,
        tahun_akademik_id: UUID | None,
        status_data: str | None,
        filter_lain: dict[str, str],
    ) -> dict[str, Any]:
        cakupan = pemanggil.cakupan_baca()
        if prodi_id and not cakupan.mengizinkan(prodi_id):
            raise AksesDitolak(alasan="di_luar_cakupan_prodi")
        stmt = self.spec.penangan.saring(self._dasar(), self.model, cakupan, prodi_id)
        if tahun_akademik_id and self.spec.kolom_tahun:
            stmt = stmt.where(getattr(self.model, self.spec.kolom_tahun) == tahun_akademik_id)
        if status_data and "status_data" in self.kolom_tabel:
            stmt = stmt.where(self.model.status_data == status_data)
        for k, v in filter_lain.items():
            if k in self.spec.filter_boleh and k in self.kolom_tabel:
                kolom = getattr(self.model, k)
                try:
                    tipe = kolom.type.python_type
                except NotImplementedError:
                    tipe = str
                try:
                    nilai: Any = v.lower() in {"1", "true", "ya"} if tipe is bool else tipe(v)
                except ValueError as exc:
                    raise ValidasiGagal(f"Nilai saringan '{k}' tidak sah.") from exc
                stmt = stmt.where(kolom == nilai)
        if param.q and self.spec.kolom_cari:
            like = f"%{param.q}%"
            stmt = stmt.where(func.coalesce(func.concat_ws(" ", *[getattr(self.model, c) for c in self.spec.kolom_cari]), "").ilike(like))
        urut = param.urutan(
            {c: getattr(self.model, c) for c in self.kolom_tabel},
            [self.model.updated_at.desc()] if "updated_at" in self.kolom_tabel else [self.model.id],
        )
        peta = {p.id: p for p in await referensi.daftar_prodi(self.session)}
        hasil = await halaman(self.session, stmt, param, urut, lambda o: o)
        hasil["data"] = [await self._ke_dict(o, peta) for o in hasil["data"]]
        return hasil

    async def rincian(self, pemanggil: Pemanggil, obj_id: UUID) -> dict[str, Any]:
        obj = await self._muat(obj_id)
        await self._wajib_baca(pemanggil, obj)
        return await self._ke_dict(obj)

    async def butir_terkait(self, pemanggil: Pemanggil, obj_id: UUID) -> list[dict[str, Any]]:
        obj = await self._muat(obj_id)
        await self._wajib_baca(pemanggil, obj)
        return await self._butir_terkait()

    async def _butir_terkait(self) -> list[dict[str, Any]]:
        return [
            {"instrumen": b.instrumen, "kode": b.kode, "elemen": b.elemen}
            for b in await instrumen.butir_terpasok_oleh(self.session, self.spec.entitas)
        ]

    # ------------------------------------------------------------ duplikasi (EC-03)
    async def _cari_duplikat(self, data: dict[str, Any], abaikan_kemiripan: bool, kecuali: UUID | None = None) -> list[dict[str, Any]]:
        sp = self.spec
        kandidat: list[dict[str, Any]] = []

        def tambah(o: Any, alasan: str) -> None:
            kandidat.append(
                {
                    "id": o.id,
                    "judul": getattr(o, "judul", None) or getattr(o, "nama", None),
                    "doi": getattr(o, "doi", None),
                    "alasan": alasan,
                }
            )

        dasar = self._dasar()
        if kecuali:
            dasar = dasar.where(self.model.id != kecuali)
        if sp.entitas == JenisEntitas.PUBLIKASI:
            if data.get("doi"):
                for o in (await self.session.execute(dasar.where(func.lower(m.Publikasi.doi) == data["doi"].lower()))).scalars():
                    tambah(o, "doi_sama")
            if not kandidat and not abaikan_kemiripan and data.get("judul"):
                judul = data["judul"].strip().lower()
                q = dasar.where(
                    (func.lower(m.Publikasi.judul) == judul) | (func.similarity(m.Publikasi.judul, data["judul"]) > 0.85)
                ).limit(5)
                for o in (await self.session.execute(q)).scalars():
                    tambah(o, "judul_mirip")
        elif sp.entitas == JenisEntitas.DOSEN:
            for kol in ("nidn", "nuptk"):
                if data.get(kol):
                    for o in (await self.session.execute(dasar.where(getattr(m.Dosen, kol) == data[kol]))).scalars():
                        tambah(o, f"{kol}_sama")
        elif sp.entitas in (JenisEntitas.PENELITIAN, JenisEntitas.PKM) and not abaikan_kemiripan and data.get("judul"):
            q = dasar.where(
                (m.KegiatanTridharma.tahun_akademik_id == data["tahun_akademik_id"])
                & (func.similarity(m.KegiatanTridharma.judul, data["judul"]) > 0.9)
            ).limit(5)
            for o in (await self.session.execute(q)).scalars():
                tambah(o, "judul_mirip")
        elif sp.entitas == JenisEntitas.KERJA_SAMA and not abaikan_kemiripan:
            q = dasar.where(
                (func.lower(m.KerjaSama.nama_mitra) == data["nama_mitra"].lower())
                & (func.lower(m.KerjaSama.judul_kegiatan) == data["judul_kegiatan"].lower())
                & (m.KerjaSama.tanggal_mulai == data["tanggal_mulai"])
            )
            for o in (await self.session.execute(q)).scalars():
                tambah(o, "kerja_sama_sama")
        return kandidat

    # ------------------------------------------------------------ tulis
    def _kolom_input(self, data: dict[str, Any]) -> dict[str, Any]:
        return {k: v for k, v in data.items() if k in self.kolom_tabel and k not in {"id", "row_version"}}

    async def _terapkan_prodi(self, pemanggil: Pemanggil, data: dict[str, Any]) -> list[tuple[UUID, Decimal]]:
        prodi = self.spec.penangan.prodi_dari_input(pemanggil, data)
        if not prodi and not pemanggil.punya(SF):
            raise ValidasiGagal("Prodi wajib ditentukan (afiliasi/atribusi prodi).")
        return prodi

    async def buat(self, pemanggil: Pemanggil, skema: BaseModel, *, abaikan_kemiripan: bool = False) -> dict[str, Any]:
        data = {k: getattr(skema, k) for k in type(skema).model_fields}
        prodi = await self._terapkan_prodi(pemanggil, data)
        self.spec.penangan.lengkapi_input(data, prodi)
        prodi_set = {p for p, _ in prodi}
        obj = self.model(**self._kolom_input(data), **self.spec.tetap)
        await self._wajib_tulis(pemanggil, prodi_set, obj)
        duplikat = await self._cari_duplikat(data, abaikan_kemiripan)
        if duplikat:
            raise Duplikat(f"{self.spec.nama} serupa sudah tercatat.", details=duplikat)
        async with self.uow:
            obj.created_by = obj.updated_by = pemanggil.id
            self.session.add(obj)
            try:
                await self.session.flush()
                await self.spec.penangan.simpan(self.session, obj, data, prodi)
                await self._simpan_dosen_terkait(obj, data)
                await self.session.flush()
            except IntegrityError as exc:
                raise self._galat_integritas(exc) from exc
            await audit.catat(
                self.session,
                f"{self.spec.path}.buat",
                pemanggil=pemanggil,
                entitas=self.spec.path,
                entitas_id=obj.id,
                prodi_id=next(iter(prodi_set)) if len(prodi_set) == 1 else None,
            )
        hasil = await self._ke_dict(obj)
        hasil["butir_terkait"] = await self._butir_terkait()
        return hasil

    def _galat_integritas(self, exc: IntegrityError) -> GalatDomain:
        pesan = str(exc.orig)
        if "uq_" in pesan or "unique" in pesan.lower() or "duplicate" in pesan.lower():
            return Duplikat(f"{self.spec.nama} dengan kunci yang sama sudah ada.")
        return ValidasiGagal("Data melanggar batasan basis data.", details=[{"pesan": pesan.splitlines()[0]}])

    async def ubah(self, pemanggil: Pemanggil, obj_id: UUID, patch: BaseModel, *, abaikan_kemiripan: bool = False) -> dict[str, Any]:
        obj = await self._muat(obj_id)
        prodi_lama = await self.spec.penangan.prodi_ids(self.session, obj)
        await self._wajib_tulis(pemanggil, prodi_lama, obj)
        if patch.row_version != obj.row_version:  # type: ignore[attr-defined]
            raise KonflikVersi(details=[{"row_version": obj.row_version, "updated_at": obj.updated_at, "data": await self._ke_dict(obj)}])
        data = {k: getattr(patch, k) for k in patch.model_fields_set if k != "row_version"}
        if not data:
            return await self._ke_dict(obj)
        if "prodi_id" in data and data["prodi_id"] != getattr(obj, "prodi_id", None):
            raise ValidasiGagal("prodi_id tidak dapat diubah.")
        prodi_baru: list[tuple[UUID, Decimal]] | None = None
        if self.bidang_ekstra & set(data) & {"atribusi_prodi", "prodi_ids"}:
            prodi_baru = self.spec.penangan.prodi_dari_input(pemanggil, data)
            if not prodi_baru:
                raise ValidasiGagal("Minimal satu prodi harus dipertahankan.")
            await self._wajib_tulis(pemanggil, {p for p, _ in prodi_baru}, obj)
        gabungan = {k: getattr(obj, k) for k in self.kolom_tabel} | self._kolom_input(data)
        if self.spec.kolom_hanya_sf and data.get(self.spec.kolom_hanya_sf) and not pemanggil.punya(SF):
            raise AksesDitolak(alasan="data_bersama_dikelola_staf_fakultas")
        duplikat = await self._cari_duplikat(gabungan, abaikan_kemiripan, kecuali=obj.id)
        if duplikat:
            raise Duplikat(f"{self.spec.nama} serupa sudah tercatat.", details=duplikat)
        async with self.uow:
            for k, v in self._kolom_input(data).items():
                setattr(obj, k, v)
            if "status_data" in self.kolom_tabel:
                obj.status_data = m.StatusData.DRAFT  # perubahan membatalkan verifikasi sebelumnya
            obj.updated_by = pemanggil.id
            try:
                await self.session.flush()
                if prodi_baru is not None:
                    await self.spec.penangan.simpan(self.session, obj, data, prodi_baru)
                await self._simpan_dosen_terkait(obj, data)
                await self.session.flush()
            except StaleDataError as exc:
                await self.session.rollback()
                raise KonflikVersi() from exc
            except IntegrityError as exc:
                raise self._galat_integritas(exc) from exc
            await audit.catat(
                self.session,
                f"{self.spec.path}.ubah",
                pemanggil=pemanggil,
                entitas=self.spec.path,
                entitas_id=obj.id,
                detail={"kolom": sorted(data)},
            )
        return await self._ke_dict(obj)

    async def hapus(self, pemanggil: Pemanggil, obj_id: UUID) -> None:
        pemanggil.wajib(K)
        obj = await self._muat(obj_id)
        prodi = await self.spec.penangan.prodi_ids(self.session, obj)
        if not any(pemanggil.cakupan_tulis().mengizinkan(p) for p in prodi):
            raise AksesDitolak(alasan="di_luar_cakupan_prodi")
        await self._wajib_tulis(pemanggil, prodi, obj)
        if "status_data" in self.kolom_tabel and obj.status_data != m.StatusData.DRAFT:
            raise KonflikStatus("Hanya data berstatus draft yang dapat dihapus (EC-09).")
        async with self.uow:
            if self.spec.soft_delete:
                obj.deleted_at = _sekarang()
            else:
                await self.session.delete(obj)
            await bukti.lepas_semua_data_inti(self.session, self.spec.entitas, obj.id)
            await audit.catat(self.session, f"{self.spec.path}.hapus", pemanggil=pemanggil, entitas=self.spec.path, entitas_id=obj_id)

    async def verifikasi(self, pemanggil: Pemanggil, obj_id: UUID) -> dict[str, Any]:
        pemanggil.wajib(K)
        obj = await self._muat(obj_id)
        prodi = await self.spec.penangan.prodi_ids(self.session, obj)
        if not any(pemanggil.cakupan_tulis().mengizinkan(p) for p in prodi):
            raise AksesDitolak(alasan="di_luar_cakupan_prodi")
        await self._wajib_tulis(pemanggil, prodi, obj)
        async with self.uow:
            obj.status_data = m.StatusData.TERVERIFIKASI
            obj.updated_by = pemanggil.id
            await audit.catat(self.session, f"{self.spec.path}.verifikasi", pemanggil=pemanggil, entitas=self.spec.path, entitas_id=obj_id)
        return await self._ke_dict(obj)

    # ------------------------------------------------------------ bukti
    async def daftar_bukti(self, pemanggil: Pemanggil, obj_id: UUID) -> list[dict[str, Any]]:
        obj = await self._muat(obj_id)
        await self._wajib_baca(pemanggil, obj)
        return await bukti.daftar_bukti_data_inti(self.session, self.spec.entitas, obj_id)

    async def lampirkan_bukti(self, pemanggil: Pemanggil, obj_id: UUID, bukti_id: UUID) -> list[dict[str, Any]]:
        obj = await self._muat(obj_id)
        await self._wajib_tulis(pemanggil, await self.spec.penangan.prodi_ids(self.session, obj), obj)
        async with self.uow:
            await bukti.tautkan_data_inti(self.session, bukti_id, self.spec.entitas, obj_id, pemanggil.id)
            await audit.catat(
                self.session,
                f"{self.spec.path}.bukti.lampir",
                pemanggil=pemanggil,
                entitas=self.spec.path,
                entitas_id=obj_id,
                detail={"bukti_id": bukti_id},
            )
        return await bukti.daftar_bukti_data_inti(self.session, self.spec.entitas, obj_id)

    async def lepas_bukti(self, pemanggil: Pemanggil, obj_id: UUID, bukti_id: UUID) -> None:
        obj = await self._muat(obj_id)
        await self._wajib_tulis(pemanggil, await self.spec.penangan.prodi_ids(self.session, obj), obj)
        async with self.uow:
            await bukti.lepas_data_inti(self.session, bukti_id, self.spec.entitas, obj_id)
            await audit.catat(
                self.session,
                f"{self.spec.path}.bukti.lepas",
                pemanggil=pemanggil,
                entitas=self.spec.path,
                entitas_id=obj_id,
                detail={"bukti_id": bukti_id},
            )

    # ------------------------------------------------------------ gabung publikasi (EC-03)
    async def gabung(self, pemanggil: Pemanggil, tujuan_id: UUID, sumber_id: UUID) -> dict[str, Any]:
        if self.spec.entitas != JenisEntitas.PUBLIKASI:
            raise TidakDitemukan()
        pemanggil.wajib(K)
        if tujuan_id == sumber_id:
            raise ValidasiGagal("Sumber dan tujuan penggabungan harus berbeda.")
        tujuan, sumber = await self._muat(tujuan_id), await self._muat(sumber_id)
        for o in (tujuan, sumber):
            await self._wajib_tulis(pemanggil, await self.spec.penangan.prodi_ids(self.session, o), o)
        async with self.uow:
            for tabel, kunci in ((m.PublikasiDosen, "dosen_id"), (m.PublikasiProdi, "prodi_id")):
                ada = set((await self.session.execute(select(getattr(tabel, kunci)).where(tabel.publikasi_id == tujuan_id))).scalars())
                for r in (await self.session.execute(select(tabel).where(tabel.publikasi_id == sumber_id))).scalars():
                    if getattr(r, kunci) not in ada:
                        self.session.add(
                            tabel(**{c.key: getattr(r, c.key) for c in inspect(tabel).mapper.column_attrs} | {"publikasi_id": tujuan_id})
                        )
                await self.session.execute(delete(tabel).where(tabel.publikasi_id == sumber_id))
            await bukti.pindahkan_data_inti(self.session, self.spec.entitas, sumber_id, tujuan_id)
            for kolom in ("doi", "url", "nama_media", "tingkat"):
                if not getattr(tujuan, kolom) and getattr(sumber, kolom):
                    setattr(tujuan, kolom, getattr(sumber, kolom))
            tujuan.jumlah_sitasi = max(tujuan.jumlah_sitasi, sumber.jumlah_sitasi)
            sumber.deleted_at = _sekarang()
            sumber.doi = None
            await audit.catat(
                self.session,
                "publikasi.gabung",
                pemanggil=pemanggil,
                entitas="publikasi",
                entitas_id=tujuan_id,
                detail={"sumber_id": sumber_id},
            )
        return await self._ke_dict(tujuan)

    # ------------------------------------------------------------ impor CSV/XLSX (EC-17)
    @staticmethod
    def _baca_berkas(nama: str, isi: bytes) -> list[dict[str, Any]]:
        if nama.lower().endswith(".csv"):
            teks = isi.decode("utf-8-sig")
            return [dict(r) for r in csv.DictReader(io.StringIO(teks))]
        try:
            ws = load_workbook(io.BytesIO(isi), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise ValidasiGagal("Berkas harus CSV atau XLSX yang valid.") from exc
        baris = list(ws.iter_rows(values_only=True))
        if not baris:
            return []
        judul = [str(c).strip() if c is not None else "" for c in baris[0]]
        return [
            {judul[i]: r[i] for i in range(min(len(judul), len(r))) if judul[i]} for r in baris[1:] if any(c not in (None, "") for c in r)
        ]

    def _siapkan_baris(
        self, baris: dict[str, Any], prodi_kode: dict[str, UUID], tahun_kode: dict[str, UUID] | None = None
    ) -> dict[str, Any]:
        hasil: dict[str, Any] = {}
        for k, v in baris.items():
            if isinstance(v, str):
                v = v.strip()
            if v in ("", None):
                continue
            if k == "tahun_akademik":  # kode yang mudah dibaca (mis. 2024/2025-Ganjil) → UUID
                tid = (tahun_kode or {}).get(str(v).strip().upper())
                if tid is None:
                    raise ValueError(f"Tahun akademik tidak dikenal: {v} (lihat lembar 'referensi' pada templat).")
                hasil["tahun_akademik_id"] = tid
                continue
            if k == "prodi_kode":
                ids = [prodi_kode[x.strip().upper()] for x in str(v).split(",") if x.strip().upper() in prodi_kode]
                if not ids:
                    raise ValueError(f"Kode prodi tidak dikenal: {v}")
                if "prodi_id" in self.spec.skema.model_fields:
                    hasil["prodi_id"] = ids[0]
                elif "prodi_ids" in self.spec.skema.model_fields:
                    hasil["prodi_ids"] = ids
                elif "atribusi_prodi" in self.spec.skema.model_fields:
                    hasil["atribusi_prodi"] = [{"prodi_id": i, "proporsi": round(1 / len(ids), 3)} for i in ids]
                continue
            medan = self.spec.skema.model_fields.get(k)
            if medan is not None and str(medan.annotation).startswith("list[") and isinstance(v, str):
                v = [x.strip() for x in v.split(",") if x.strip()]
            hasil[k] = v
        return hasil

    async def templat_impor(self, format: str) -> tuple[bytes, str, str]:
        """Templat impor untuk sumber ini: (isi, tipe media, nama berkas)."""
        daftar_prodi = await referensi.daftar_prodi(self.session)
        daftar_tahun = await referensi.daftar_tahun_akademik(self.session)
        aktif = next((t for t in daftar_tahun if t.is_aktif), daftar_tahun[0] if daftar_tahun else None)
        kolom = kolom_impor(
            self.spec,
            kode_prodi=daftar_prodi[0].kode if daftar_prodi else "KODE-PRODI",
            kode_tahun=aktif.kode if aktif else "2025/2026",
        )
        dasar = f"templat-impor-{self.spec.path}"
        if format == "csv":
            return buat_csv(kolom), TIPE_CSV, f"{dasar}.csv"
        referensi_lembar = {
            "prodi": (
                ["kode (isi di kolom prodi_kode)", "nama", "jenjang", "LAM"],
                [[p.kode, p.nama, p.jenjang, p.lam_kode] for p in daftar_prodi],
            ),
            "tahun_akademik": (
                ["kode (isi di kolom tahun_akademik)", "mulai", "selesai", "aktif"],
                [[t.kode, t.tanggal_mulai.isoformat(), t.tanggal_selesai.isoformat(), "ya" if t.is_aktif else ""] for t in daftar_tahun],
            ),
        }
        return (
            buat_xlsx(self.spec.nama, kolom, catatan=list(self.spec.catatan_impor), referensi=referensi_lembar),
            TIPE_XLSX,
            f"{dasar}.xlsx",
        )

    async def impor(self, pemanggil: Pemanggil, nama: str, isi: bytes, *, atomik: bool = False) -> dict[str, Any]:
        baris_mentah = self._baca_berkas(nama, isi)
        if len(baris_mentah) > 2000:
            raise ValidasiGagal("Maksimal 2000 baris per impor.")
        prodi_kode = {p.kode.upper(): p.id for p in await referensi.daftar_prodi(self.session)}
        tahun_kode = {t.kode.strip().upper(): t.id for t in await referensi.daftar_tahun_akademik(self.session)}
        berhasil, gagal = 0, []
        for no, mentah in enumerate(baris_mentah, start=2):  # baris 1 = judul
            try:
                skema = self.spec.skema.model_validate(self._siapkan_baris(mentah, prodi_kode, tahun_kode))
                async with self.session.begin_nested():
                    await self._buat_tanpa_komit(pemanggil, skema)
                berhasil += 1
            except ValidationError as exc:
                for e in exc.errors():
                    gagal.append({"baris": no, "kolom": ".".join(str(x) for x in e["loc"]), "pesan": e["msg"]})
            except (GalatDomain, ValueError) as exc:
                gagal.append({"baris": no, "kolom": "-", "pesan": getattr(exc, "pesan", str(exc))})
        if gagal and atomik:
            await self.session.rollback()
            return {"status": "dibatalkan", "baris_berhasil": 0, "baris_gagal": gagal}
        await audit.catat(
            self.session,
            f"{self.spec.path}.impor",
            pemanggil=pemanggil,
            entitas=self.spec.path,
            detail={"berhasil": berhasil, "gagal": len(gagal)},
        )
        await self.session.commit()
        return {"status": "selesai", "baris_berhasil": berhasil, "baris_gagal": gagal}

    async def _buat_tanpa_komit(self, pemanggil: Pemanggil, skema: BaseModel) -> None:
        data = {k: getattr(skema, k) for k in type(skema).model_fields}
        prodi = await self._terapkan_prodi(pemanggil, data)
        self.spec.penangan.lengkapi_input(data, prodi)
        prodi_set = {p for p, _ in prodi}
        obj = self.model(**self._kolom_input(data), **self.spec.tetap)
        await self._wajib_tulis(pemanggil, prodi_set, obj)
        duplikat = await self._cari_duplikat(data, abaikan_kemiripan=False)
        if duplikat:
            raise Duplikat(f"{self.spec.nama} serupa sudah tercatat.", details=duplikat)
        obj.created_by = obj.updated_by = pemanggil.id
        self.session.add(obj)
        try:
            await self.session.flush()
            await self.spec.penangan.simpan(self.session, obj, data, prodi)
            await self._simpan_dosen_terkait(obj, data)
            await self.session.flush()
        except IntegrityError as exc:
            raise self._galat_integritas(exc) from exc
