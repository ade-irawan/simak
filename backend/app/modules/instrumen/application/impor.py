"""Impor versi instrumen dari XLSX (sebagai draf) dan pembuatan templat impor.

Lembar yang dibaca (baris pertama = judul kolom, huruf kecil):
  kriteria  : kode, nama, bobot, urutan
  butir     : kriteria, kode, nomor_urut, elemen, deskriptor, jenis_ipo, tahap_ppepp, jenis_penilaian, bobot,
              panduan, skema_narasi (JSON)
  rubrik    : butir, skor, deskriptor
  peringkat : kode, nama, tingkat, masa_berlaku_tahun, nilai_akhir_min, nilai_akhir_maks, syarat (JSON)
  ambang    : butir, peringkat, skor_min, is_syarat_perlu
"""

from __future__ import annotations

import io
import json
from decimal import Decimal, InvalidOperation
from typing import Any
from uuid import UUID

from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import ValidasiGagal
from app.modules.instrumen.domain.entities import JenisIpo, JenisPenilaian, StatusInstrumen, TahapPpepp
from app.modules.instrumen.infrastructure.models import (
    InstrumenButir,
    InstrumenButirAmbang,
    InstrumenButirRubrik,
    InstrumenKriteria,
    InstrumenPeringkat,
    InstrumenVersi,
)

KOLOM: dict[str, list[str]] = {
    "kriteria": ["kode", "nama", "bobot", "urutan"],
    "butir": [
        "kriteria",
        "kode",
        "nomor_urut",
        "elemen",
        "deskriptor",
        "jenis_ipo",
        "tahap_ppepp",
        "jenis_penilaian",
        "bobot",
        "panduan",
        "skema_narasi",
    ],
    "rubrik": ["butir", "skor", "deskriptor"],
    "peringkat": ["kode", "nama", "tingkat", "masa_berlaku_tahun", "nilai_akhir_min", "nilai_akhir_maks", "syarat"],
    "ambang": ["butir", "peringkat", "skor_min", "is_syarat_perlu"],
}


def buat_templat_xlsx() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for nama, kolom in KOLOM.items():
        ws = wb.create_sheet(nama)
        ws.append(kolom)
    wb["kriteria"].append(["A", "Contoh Kriteria", None, 1])
    wb["butir"].append(
        ["A", "1", 1, "Contoh elemen penilaian", None, None, None, "kualitatif", None, None, '[{"bagian":"utama","judul":"Uraian"}]']
    )
    wb["rubrik"].append(["1", 4, "Deskriptor skor 4"])
    wb["peringkat"].append(["terakreditasi", "Terakreditasi", 1, None, None, None, "{}"])
    wb["ambang"].append(["1", "terakreditasi", 3, "TRUE"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _baca_lembar(wb: Any, nama: str) -> list[tuple[int, dict[str, Any]]]:
    if nama not in wb.sheetnames:
        return []
    ws = wb[nama]
    baris = list(ws.iter_rows(values_only=True))
    if not baris:
        return []
    judul = [str(c).strip().lower() if c is not None else "" for c in baris[0]]
    hasil = []
    for no, r in enumerate(baris[1:], start=2):
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        hasil.append((no, {judul[i]: r[i] for i in range(min(len(judul), len(r))) if judul[i]}))
    return hasil


def _desimal(nilai: Any, lembar: str, no: int, kolom: str, galat: list[dict[str, Any]]) -> Decimal | None:
    if nilai is None or str(nilai).strip() == "":
        return None
    try:
        return Decimal(str(nilai).replace(",", "."))
    except InvalidOperation:
        galat.append({"lembar": lembar, "baris": no, "kolom": kolom, "pesan": "Bukan angka."})
        return None


def _json(nilai: Any, lembar: str, no: int, kolom: str, bawaan: Any, galat: list[dict[str, Any]]) -> Any:
    if nilai is None or str(nilai).strip() == "":
        return bawaan
    try:
        return json.loads(str(nilai))
    except json.JSONDecodeError:
        galat.append({"lembar": lembar, "baris": no, "kolom": kolom, "pesan": "JSON tidak valid."})
        return bawaan


def _enum(kelas: Any, nilai: Any, lembar: str, no: int, kolom: str, galat: list[dict[str, Any]]) -> Any:
    if nilai is None or str(nilai).strip() == "":
        return None
    try:
        return kelas(str(nilai).strip().lower())
    except ValueError:
        galat.append({"lembar": lembar, "baris": no, "kolom": kolom, "pesan": f"Nilai harus salah satu dari {[m.value for m in kelas]}."})
        return None


async def impor_instrumen(
    session: AsyncSession,
    isi: bytes,
    *,
    lam_id: UUID,
    kode_versi: str,
    nama: str,
    jenjang: str,
    skor_maks: int,
    label_skor: dict[str, str],
    nilai_akhir_maks: Decimal | None,
    dasar_hukum: str | None,
    disusun_oleh: UUID,
) -> dict[str, Any]:
    try:
        wb = load_workbook(io.BytesIO(isi), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidasiGagal("Berkas bukan XLSX yang valid.") from exc

    galat: list[dict[str, Any]] = []
    kriteria_baris = _baca_lembar(wb, "kriteria")
    butir_baris = _baca_lembar(wb, "butir")
    if not kriteria_baris or not butir_baris:
        raise ValidasiGagal("Lembar 'kriteria' dan 'butir' wajib berisi data.")

    versi = InstrumenVersi(
        lam_id=lam_id,
        kode_versi=kode_versi,
        nama=nama,
        jenjang=jenjang,
        dasar_hukum=dasar_hukum,
        status=StatusInstrumen.DRAF,
        skor_maks=skor_maks,
        label_skor=label_skor,
        nilai_akhir_maks=nilai_akhir_maks,
        disusun_oleh=disusun_oleh,
    )
    session.add(versi)
    await session.flush()

    kriteria: dict[str, InstrumenKriteria] = {}
    for no, r in kriteria_baris:
        kode = str(r.get("kode") or "").strip()
        if not kode or not r.get("nama"):
            galat.append({"lembar": "kriteria", "baris": no, "kolom": "kode/nama", "pesan": "Wajib diisi."})
            continue
        urutan = int(r.get("urutan") or len(kriteria) + 1)
        k = InstrumenKriteria(
            instrumen_versi_id=versi.id,
            kode=kode,
            nama=str(r["nama"]),
            bobot=_desimal(r.get("bobot"), "kriteria", no, "bobot", galat),
            urutan=urutan,
        )
        kriteria[kode] = k
        session.add(k)
    await session.flush()

    butir: dict[str, InstrumenButir] = {}
    for no, r in butir_baris:
        kode = str(r.get("kode") or "").strip()
        k = kriteria.get(str(r.get("kriteria") or "").strip())
        if not kode or not r.get("elemen") or k is None:
            galat.append(
                {
                    "lembar": "butir",
                    "baris": no,
                    "kolom": "kriteria/kode/elemen",
                    "pesan": "Kode, elemen, dan kriteria yang dikenal wajib diisi.",
                }
            )
            continue
        if kode in butir:
            galat.append({"lembar": "butir", "baris": no, "kolom": "kode", "pesan": f"Kode butir ganda: {kode}."})
            continue
        b = InstrumenButir(
            instrumen_versi_id=versi.id,
            kriteria_id=k.id,
            kode=kode,
            nomor_urut=int(r.get("nomor_urut") or len(butir) + 1),
            elemen=str(r["elemen"]),
            deskriptor=r.get("deskriptor"),
            panduan=r.get("panduan"),
            jenis_ipo=_enum(JenisIpo, r.get("jenis_ipo"), "butir", no, "jenis_ipo", galat),
            tahap_ppepp=_enum(TahapPpepp, r.get("tahap_ppepp"), "butir", no, "tahap_ppepp", galat),
            jenis_penilaian=_enum(JenisPenilaian, r.get("jenis_penilaian"), "butir", no, "jenis_penilaian", galat)
            or JenisPenilaian.KUALITATIF,
            bobot=_desimal(r.get("bobot"), "butir", no, "bobot", galat),
            skema_narasi=_json(r.get("skema_narasi"), "butir", no, "skema_narasi", [], galat),
        )
        butir[kode] = b
        session.add(b)
    await session.flush()

    n_rubrik = 0
    for no, r in _baca_lembar(wb, "rubrik"):
        b = butir.get(str(r.get("butir") or "").strip())
        if b is None or r.get("skor") is None or not r.get("deskriptor"):
            galat.append(
                {"lembar": "rubrik", "baris": no, "kolom": "butir/skor/deskriptor", "pesan": "Data tidak lengkap atau butir tidak dikenal."}
            )
            continue
        session.add(InstrumenButirRubrik(instrumen_butir_id=b.id, skor=int(r["skor"]), deskriptor=str(r["deskriptor"])))
        n_rubrik += 1

    peringkat: dict[str, InstrumenPeringkat] = {}
    for no, r in _baca_lembar(wb, "peringkat"):
        kode = str(r.get("kode") or "").strip()
        if not kode or not r.get("nama") or r.get("tingkat") is None:
            galat.append({"lembar": "peringkat", "baris": no, "kolom": "kode/nama/tingkat", "pesan": "Wajib diisi."})
            continue
        p = InstrumenPeringkat(
            instrumen_versi_id=versi.id,
            kode=kode,
            nama=str(r["nama"]),
            tingkat=int(r["tingkat"]),
            masa_berlaku_tahun=int(r["masa_berlaku_tahun"]) if r.get("masa_berlaku_tahun") not in (None, "") else None,
            nilai_akhir_min=_desimal(r.get("nilai_akhir_min"), "peringkat", no, "nilai_akhir_min", galat),
            nilai_akhir_maks=_desimal(r.get("nilai_akhir_maks"), "peringkat", no, "nilai_akhir_maks", galat),
            syarat=_json(r.get("syarat"), "peringkat", no, "syarat", {}, galat),
        )
        peringkat[kode] = p
        session.add(p)
    await session.flush()

    n_ambang = 0
    for no, r in _baca_lembar(wb, "ambang"):
        b = butir.get(str(r.get("butir") or "").strip())
        p = peringkat.get(str(r.get("peringkat") or "").strip())
        skor_min = _desimal(r.get("skor_min"), "ambang", no, "skor_min", galat)
        if b is None or p is None or skor_min is None:
            galat.append(
                {"lembar": "ambang", "baris": no, "kolom": "butir/peringkat/skor_min", "pesan": "Data tidak lengkap atau tidak dikenal."}
            )
            continue
        session.add(
            InstrumenButirAmbang(
                instrumen_butir_id=b.id,
                peringkat_id=p.id,
                skor_min=skor_min,
                is_syarat_perlu=str(r.get("is_syarat_perlu")).strip().lower() in {"true", "1", "ya", "y"},
            )
        )
        n_ambang += 1

    if galat:
        await session.rollback()
        raise ValidasiGagal("Impor instrumen dibatalkan; perbaiki baris bermasalah.", details=galat[:200])
    await session.flush()
    return {
        "id": versi.id,
        "status": "selesai",
        "kriteria": len(kriteria),
        "butir": len(butir),
        "rubrik": n_rubrik,
        "peringkat": len(peringkat),
        "ambang": n_ambang,
    }
