"""Use case dokumen eksternal: SIMAK menerima, mengarsipkan, dan memetakan hasil — bukan menyelenggarakan survei."""

from __future__ import annotations

import csv
import io
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import delete, select
from sqlalchemy.dialects.postgresql import insert as pg_insert

from app.core import audit
from app.core.exceptions import AksesDitolak, TidakDitemukan, ValidasiGagal
from app.core.pagination import ParamHalaman, halaman
from app.core.permissions import SF, SP, D, K, Pemanggil
from app.core.unit_of_work import UnitOfWork
from app.infrastructure.documents.templat_impor import TIPE_CSV, TIPE_XLSX, KolomImpor, buat_csv, buat_xlsx
from app.modules.bukti import public as bukti
from app.modules.dokumen_eksternal.api.schemas import IndikatorIn
from app.modules.dokumen_eksternal.infrastructure.models import (
    DokumenEksternal,
    DokumenEksternalProdi,
    IndikatorEksternal,
)
from app.modules.instrumen import public as instrumen
from app.modules.referensi import public as referensi
from app.modules.siklus import public as siklus

MIN_RESPONDEN_BAWAAN = 30


class LayananDokumenEksternal:
    def __init__(self, uow: UnitOfWork):
        self.uow, self.session = uow, uow.session

    async def _prodi_dokumen(self, dokumen_id: UUID) -> set[UUID]:
        return set(
            (
                await self.session.execute(select(DokumenEksternalProdi.prodi_id).where(DokumenEksternalProdi.dokumen_id == dokumen_id))
            ).scalars()
        )

    async def _ke_dict(self, d: DokumenEksternal, prodi: set[UUID] | None = None) -> dict[str, Any]:
        peta = {p.id: p for p in await referensi.daftar_prodi(self.session)}
        prodi = prodi if prodi is not None else await self._prodi_dokumen(d.id)
        ringkas = (await bukti.ringkas_bukti(self.session, [d.bukti_id])).get(d.bukti_id)
        return {
            "id": d.id,
            "jenis": d.jenis,
            "judul": d.judul,
            "sumber_unit": d.sumber_unit,
            "tahun_akademik_id": d.tahun_akademik_id,
            "periode_mulai": d.periode_mulai,
            "periode_selesai": d.periode_selesai,
            "populasi": d.populasi,
            "jumlah_responden": d.jumlah_responden,
            "metodologi": d.metodologi,
            "diterima_pada": d.diterima_pada,
            "bukti": ringkas,
            "prodi": [{"id": p, "kode": peta[p].kode, "nama": peta[p].nama} for p in prodi if p in peta],
            "created_at": d.created_at,
        }

    async def _muat(self, dokumen_id: UUID) -> DokumenEksternal:
        d = (
            await self.session.execute(
                select(DokumenEksternal).where(DokumenEksternal.id == dokumen_id, DokumenEksternal.deleted_at.is_(None))
            )
        ).scalar_one_or_none()
        if d is None:
            raise TidakDitemukan("Dokumen tidak ditemukan.")
        return d

    async def _wajib_baca(self, pemanggil: Pemanggil, prodi: set[UUID]) -> None:
        cakupan = pemanggil.cakupan_baca()
        if not cakupan.semua and not any(cakupan.mengizinkan(p) for p in prodi):
            raise AksesDitolak(alasan="di_luar_cakupan_prodi")

    async def _sinkron_eksternal(self, prodi_ids: set[UUID]) -> int:
        """Mencabut/memasang tanda 'Menunggu dokumen dari Direktorat' (EC-15). Mengembalikan jumlah butir dibebaskan."""
        butir = [b.butir_id for b in await instrumen.butir_terpasok_oleh(self.session, instrumen.JenisEntitas.INDIKATOR_EKSTERNAL)]
        dibebaskan = 0
        from app.modules.dokumen_eksternal.public import prodi_punya_indikator

        for p in prodi_ids:
            dibebaskan += await siklus.sinkron_menunggu_eksternal(self.session, p, await prodi_punya_indikator(self.session, p), butir)
        return dibebaskan

    # ------------------------------------------------------------ baca
    async def daftar(
        self, pemanggil: Pemanggil, param: ParamHalaman, *, jenis: str | None, prodi_id: UUID | None, tahun_akademik_id: UUID | None
    ) -> dict[str, Any]:
        cakupan = pemanggil.cakupan_baca()
        stmt = select(DokumenEksternal).where(DokumenEksternal.deleted_at.is_(None))
        if prodi_id and not cakupan.mengizinkan(prodi_id):
            raise AksesDitolak(alasan="di_luar_cakupan_prodi")
        sub = select(DokumenEksternalProdi.dokumen_id)
        if prodi_id:
            stmt = stmt.where(DokumenEksternal.id.in_(sub.where(DokumenEksternalProdi.prodi_id == prodi_id)))
        elif not cakupan.semua:
            stmt = stmt.where(DokumenEksternal.id.in_(sub.where(DokumenEksternalProdi.prodi_id.in_(cakupan.prodi_ids))))
        if jenis:
            stmt = stmt.where(DokumenEksternal.jenis == jenis)
        if tahun_akademik_id:
            stmt = stmt.where(DokumenEksternal.tahun_akademik_id == tahun_akademik_id)
        if param.q:
            stmt = stmt.where(DokumenEksternal.judul.ilike(f"%{param.q}%"))
        hasil = await halaman(self.session, stmt, param, [DokumenEksternal.created_at.desc()], lambda d: d)
        hasil["data"] = [await self._ke_dict(d) for d in hasil["data"]]
        return hasil

    async def indikator_prodi(self, pemanggil: Pemanggil, prodi_id: UUID) -> list[dict[str, Any]]:
        pemanggil.wajib_baca_prodi(prodi_id)
        baris = (
            await self.session.execute(
                select(IndikatorEksternal, DokumenEksternal.jenis, DokumenEksternal.judul, DokumenEksternal.tahun_akademik_id)
                .join(DokumenEksternal, DokumenEksternal.id == IndikatorEksternal.dokumen_id)
                .where(IndikatorEksternal.prodi_id == prodi_id, DokumenEksternal.deleted_at.is_(None))
                .order_by(DokumenEksternal.tahun_akademik_id.desc(), IndikatorEksternal.kode)
            )
        ).all()
        return [
            {
                "id": i.id,
                "dokumen_id": i.dokumen_id,
                "dokumen": judul,
                "jenis_dokumen": jenis,
                "tahun_akademik_id": ta,
                "kode": i.kode,
                "nama": i.nama,
                # EC-16: populasi/responden terlalu kecil → jangan tampilkan angka yang menyesatkan.
                "nilai": i.nilai if i.data_memadai else None,
                "satuan": i.satuan,
                "jumlah_responden": i.jumlah_responden,
                "data_memadai": i.data_memadai,
                "catatan": i.catatan,
            }
            for i, jenis, judul, ta in baris
        ]

    # ------------------------------------------------------------ tulis (Staf Fakultas)
    async def buat(self, pemanggil: Pemanggil, data: Any) -> dict[str, Any]:
        pemanggil.wajib(SF)
        await bukti.pastikan_bukti_ada(self.session, data.bukti_id)
        known = {p.id for p in await referensi.daftar_prodi(self.session)}
        if not set(data.prodi_id) <= known:
            raise ValidasiGagal("Ada prodi yang tidak dikenal.")
        async with self.uow:
            d = DokumenEksternal(**data.model_dump(exclude={"prodi_id"}), created_by=pemanggil.id, updated_by=pemanggil.id)
            self.session.add(d)
            await self.session.flush()
            for p in set(data.prodi_id):
                self.session.add(DokumenEksternalProdi(dokumen_id=d.id, prodi_id=p))
            await audit.catat(
                self.session,
                "dokumen_eksternal.buat",
                pemanggil=pemanggil,
                entitas="dokumen_eksternal",
                entitas_id=d.id,
                detail={"jenis": data.jenis},
            )
        # Dokumen tiba tetapi indikator belum ada → butir tetap menunggu hingga indikator diisi (EC-15).
        hasil = await self._ke_dict(d, set(data.prodi_id))
        hasil["butir_dibebaskan"] = 0
        return hasil

    async def ubah(self, pemanggil: Pemanggil, dokumen_id: UUID, patch: Any) -> dict[str, Any]:
        pemanggil.wajib(SF)
        d = await self._muat(dokumen_id)
        data = patch.model_dump(exclude_unset=True)
        prodi_baru = data.pop("prodi_id", None)
        async with self.uow:
            for k, v in data.items():
                setattr(d, k, v)
            if prodi_baru is not None:
                await self.session.execute(delete(DokumenEksternalProdi).where(DokumenEksternalProdi.dokumen_id == d.id))
                for p in set(prodi_baru):
                    self.session.add(DokumenEksternalProdi(dokumen_id=d.id, prodi_id=p))
            d.updated_by = pemanggil.id
            await audit.catat(
                self.session,
                "dokumen_eksternal.ubah",
                pemanggil=pemanggil,
                entitas="dokumen_eksternal",
                entitas_id=d.id,
                detail={"kolom": sorted(data)},
            )
        return await self._ke_dict(d)

    async def _min_responden(self) -> int:
        return int(await referensi.ambil_pengaturan(self.session, "eksternal.min_responden", MIN_RESPONDEN_BAWAAN))

    async def _simpan_indikator(self, d: DokumenEksternal, prodi_dokumen: set[UUID], baris: list[IndikatorIn]) -> None:
        minimal = await self._min_responden()
        for b in baris:
            if b.prodi_id not in prodi_dokumen:
                raise ValidasiGagal("Prodi tidak termasuk dalam dokumen ini.", details=[{"prodi_id": str(b.prodi_id)}])
            responden = b.jumlah_responden if b.jumlah_responden is not None else d.jumlah_responden
            memadai = responden is None or responden >= minimal
            nilai = {
                "nama": b.nama,
                "nilai": b.nilai,
                "satuan": b.satuan,
                "jumlah_responden": b.jumlah_responden,
                "data_memadai": memadai,
                "catatan": b.catatan,
            }
            await self.session.execute(
                pg_insert(IndikatorEksternal)
                .values(dokumen_id=d.id, prodi_id=b.prodi_id, kode=b.kode, **nilai)
                .on_conflict_do_update(index_elements=["dokumen_id", "prodi_id", "kode"], set_=nilai)
            )

    async def isi_indikator(self, pemanggil: Pemanggil, dokumen_id: UUID, baris: list[IndikatorIn]) -> dict[str, Any]:
        pemanggil.wajib(SF)
        d = await self._muat(dokumen_id)
        prodi = await self._prodi_dokumen(d.id)
        async with self.uow:
            await self._simpan_indikator(d, prodi, baris)
            dibebaskan = await self._sinkron_eksternal(prodi)
            await audit.catat(
                self.session,
                "dokumen_eksternal.indikator",
                pemanggil=pemanggil,
                entitas="dokumen_eksternal",
                entitas_id=d.id,
                detail={"jumlah": len(baris)},
            )
        return {"id": d.id, "baris_berhasil": len(baris), "butir_dibebaskan": dibebaskan}

    async def templat_indikator(self, format: str) -> tuple[bytes, str, str]:
        """Templat impor indikator eksternal: (isi, tipe media, nama berkas)."""
        daftar_prodi = await referensi.daftar_prodi(self.session)
        kolom = [
            KolomImpor(
                "prodi_kode",
                True,
                "kode prodi",
                "Kode prodi (lihat lembar 'referensi'). Satu baris per prodi per indikator.",
                daftar_prodi[0].kode if daftar_prodi else "KODE-PRODI",
            ),
            KolomImpor(
                "kode",
                True,
                "teks",
                "Kode indikator tanpa spasi, minimal 2 karakter. Kode yang sama pada prodi dan dokumen yang sama memperbarui nilai lama.",
                "masa_tunggu_bulan",
            ),
            KolomImpor("nama", True, "teks", "Nama indikator yang tampil di aplikasi, minimal 2 karakter.", "Masa tunggu kerja pertama"),
            KolomImpor("nilai", False, "angka (desimal pakai titik)", "Kosongkan bila angkanya memang tidak tersedia.", 3.5),
            KolomImpor("satuan", False, "teks", "Mis. bulan, persen, skor.", "bulan"),
            KolomImpor(
                "jumlah_responden",
                False,
                "bilangan bulat",
                "Bila kurang dari 30, indikator ditampilkan sebagai 'data tidak memadai'. ≥ 0.",
                120,
            ),
            KolomImpor("catatan", False, "teks", "Sumber atau keterangan tambahan.", "Tracer study 2024"),
        ]
        if format == "csv":
            return buat_csv(kolom), TIPE_CSV, "templat-impor-indikator.csv"
        isi = buat_xlsx(
            "Indikator dokumen eksternal",
            kolom,
            referensi={
                "prodi": (
                    ["kode (isi di kolom prodi_kode)", "nama", "jenjang", "LAM"],
                    [[p.kode, p.nama, p.jenjang, p.lam_kode] for p in daftar_prodi],
                )
            },
        )
        return isi, TIPE_XLSX, "templat-impor-indikator.xlsx"

    async def impor_indikator(self, pemanggil: Pemanggil, dokumen_id: UUID, nama: str, isi: bytes, atomik: bool) -> dict[str, Any]:
        """Kolom: prodi_kode, kode, nama, nilai, satuan, jumlah_responden, catatan. Baris bermasalah dilaporkan
        dengan nomornya tanpa membatalkan baris sah, kecuali ``atomik`` (EC-17)."""
        pemanggil.wajib(SF)
        d = await self._muat(dokumen_id)
        prodi = await self._prodi_dokumen(d.id)
        kode_ke_id = {p.kode.upper(): p.id for p in await referensi.daftar_prodi(self.session)}
        if nama.lower().endswith(".csv"):
            mentah = [dict(r) for r in csv.DictReader(io.StringIO(isi.decode("utf-8-sig")))]
        else:
            try:
                ws = load_workbook(io.BytesIO(isi), read_only=True, data_only=True).worksheets[0]
            except Exception as exc:
                raise ValidasiGagal("Berkas harus CSV atau XLSX yang valid.") from exc
            baris_x = list(ws.iter_rows(values_only=True))
            judul = [str(c).strip() if c is not None else "" for c in (baris_x[0] if baris_x else [])]
            mentah = [
                {judul[i]: r[i] for i in range(min(len(judul), len(r))) if judul[i]}
                for r in baris_x[1:]
                if any(c not in (None, "") for c in r)
            ]
        valid: list[IndikatorIn] = []
        gagal: list[dict[str, Any]] = []
        for no, r in enumerate(mentah, start=2):
            try:
                kode_prodi = str(r.get("prodi_kode") or "").strip().upper()
                if kode_prodi not in kode_ke_id:
                    gagal.append({"baris": no, "kolom": "prodi_kode", "pesan": f"Kode prodi tidak dikenal: {kode_prodi or '-'}."})
                    continue
                valid.append(
                    IndikatorIn.model_validate(
                        {
                            "prodi_id": kode_ke_id[kode_prodi],
                            "kode": str(r.get("kode") or "").strip(),
                            "nama": str(r.get("nama") or "").strip(),
                            "nilai": r.get("nilai") if r.get("nilai") not in ("", None) else None,
                            "satuan": r.get("satuan") or None,
                            "jumlah_responden": r.get("jumlah_responden") if r.get("jumlah_responden") not in ("", None) else None,
                            "catatan": r.get("catatan") or None,
                        }
                    )
                )
            except ValidationError as exc:
                for e in exc.errors():
                    kolom = ".".join(str(x) for x in e["loc"])
                    gagal.append(
                        {"baris": no, "kolom": kolom, "pesan": "Bukan angka." if "decimal" in e["type"] or "int" in e["type"] else e["msg"]}
                    )
        if gagal and atomik:
            return {"status": "dibatalkan", "baris_berhasil": 0, "baris_gagal": gagal}
        async with self.uow:
            ok = [v for v in valid if v.prodi_id in prodi]
            gagal += [
                {"baris": 0, "kolom": "prodi_kode", "pesan": "Prodi tidak termasuk dalam dokumen ini."}
                for v in valid
                if v.prodi_id not in prodi
            ]
            await self._simpan_indikator(d, prodi, ok)
            dibebaskan = await self._sinkron_eksternal(prodi)
            await audit.catat(
                self.session,
                "dokumen_eksternal.impor",
                pemanggil=pemanggil,
                entitas="dokumen_eksternal",
                entitas_id=d.id,
                detail={"berhasil": len(ok), "gagal": len(gagal)},
            )
        return {"id": d.id, "status": "selesai", "baris_berhasil": len(ok), "baris_gagal": gagal, "butir_dibebaskan": dibebaskan}


__all__ = ["LayananDokumenEksternal", "D", "K", "SP"]
