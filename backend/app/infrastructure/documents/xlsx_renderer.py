from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.infrastructure.documents.umum import CATATAN_ESTIMASI, judul_instrumen, waktu_lokal

JUDUL = PatternFill("solid", fgColor="0060A8")


def _tabel(ws: Any, judul: list[str], baris: list[list[Any]]) -> None:
    ws.append(judul)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), JUDUL
    for b in baris:
        ws.append(b)
    for i, kolom in enumerate(ws.columns, start=1):
        panjang = max((len(str(c.value)) for c in kolom if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(panjang + 2, 10), 70)
    ws.freeze_panes = "A2"
    for baris_ws in ws.iter_rows(min_row=2):
        for c in baris_ws:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def render_kuantitatif_xlsx(cuplikan: dict[str, Any], *, final: bool, metadata: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    skor = cuplikan.get("skor") or {}
    for k, v in [
        ("Prodi", cuplikan["prodi"]["nama"]),
        ("Instrumen", judul_instrumen(cuplikan)),
        ("Siklus", cuplikan["siklus"]["nama"]),
        ("Data per", waktu_lokal(metadata.get("data_per"))),
        ("Status", "FINAL" if final else "DRAF"),
        ("Nilai akhir estimasi", skor.get("nilai_akhir")),
        ("Prediksi peringkat", (skor.get("prediksi_peringkat") or {}).get("nama")),
        ("Catatan", CATATAN_ESTIMASI),
    ]:
        ws.append([k, v])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 70

    butir = wb.create_sheet("Butir")
    _tabel(
        butir,
        ["Kriteria", "Kode", "Elemen", "Bobot", "Jenis", "Tahap PPEPP", "Skor estimasi", "Status"],
        [
            [k["kode"], b["kode"], b["elemen"], b["bobot"], b["jenis_ipo"], b["tahap_ppepp"], b["skor_estimasi"], b["status_kelengkapan"]]
            for k in cuplikan["kriteria"]
            for b in k["butir"]
        ],
    )
    data = wb.create_sheet("Data terpetakan")
    _tabel(
        data,
        ["Kode butir", "Sumber", "Jumlah catatan", "Nilai agregat", "Keterangan"],
        [
            [b["kode"], d["sumber"], d["jumlah"], d.get("nilai"), d.get("keterangan")]
            for k in cuplikan["kriteria"]
            for b in k["butir"]
            for d in b["data_terpetakan"]
        ],
    )
    bukti = wb.create_sheet("Bukti")
    _tabel(
        bukti,
        ["Kode butir", "Judul", "Jenis", "Lokasi", "Status keterjangkauan"],
        [
            [b["kode"], x["judul"], x["jenis"], x.get("url") or x.get("nama_berkas"), x["status_cek"]]
            for k in cuplikan["kriteria"]
            for b in k["butir"]
            for x in b["bukti"]
        ],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_gap_xlsx(gap: dict[str, Any], *, metadata: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Gap analysis"
    ws.append(["Target peringkat", gap.get("target_peringkat")])
    ws.append(["Data per", waktu_lokal(metadata.get("data_per"))])
    ws.append(["Catatan", CATATAN_ESTIMASI])
    ws.append([])
    rows = [
        [b["kriteria"], b["kode"], b["bobot"], b["skor_estimasi"], b["skor_target"], b["selisih"], "Ya" if b["is_syarat_perlu"] else ""]
        for b in gap.get("butir", [])
    ]
    start = ws.max_row + 1
    ws.append(["Kriteria", "Kode", "Bobot", "Skor estimasi", "Skor target", "Selisih", "Syarat perlu"])
    for c in ws[start]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), JUDUL
    for r in rows:
        ws.append(r)
    for i in range(1, 8):
        ws.column_dimensions[get_column_letter(i)].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
