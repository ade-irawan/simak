from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.infrastructure.documents.umum import CATATAN_ESTIMASI, judul_instrumen, teks_bukti, teks_data, varian, waktu_lokal


def render_docx(cuplikan: dict[str, Any], *, final: bool, metadata: dict[str, Any]) -> bytes:
    doc = Document()
    gaya = doc.styles["Normal"]
    gaya.font.name, gaya.font.size = "Calibri", Pt(11)

    if not final:
        p = doc.sections[0].header.paragraphs[0]
        r = p.add_run("DRAF — bukan dokumen final")
        r.bold, r.font.color.rgb = True, RGBColor(0xC0, 0x00, 0x20)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    prodi, siklus = cuplikan["prodi"], cuplikan["siklus"]
    infokom = varian(cuplikan) == "lam_infokom_2_1"
    doc.add_heading("Laporan Evaluasi Diri", 0)
    doc.add_paragraph(f"Program Studi {prodi['nama']} ({prodi['jenjang']})")
    doc.add_paragraph(judul_instrumen(cuplikan))
    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Light List Accent 1"
    for k, v in [
        ("Siklus", siklus["nama"]),
        ("Tahun akademik TS", siklus["ts"]),
        ("Data per", waktu_lokal(metadata.get("data_per"))),
        (
            "Versi laporan",
            f"Pengajuan #{metadata['pengajuan_nomor']}" if metadata.get("pengajuan_nomor") else "Data berjalan (belum diajukan)",
        ),
        ("Status dokumen", "FINAL (disetujui Dekan)" if final else "DRAF"),
    ]:
        c = meta.add_row().cells
        c[0].text, c[1].text = k, str(v)

    skor = cuplikan.get("skor") or {}
    doc.add_heading("Ringkasan Skor Estimasi", 1)
    if skor.get("nilai_akhir") is not None:
        doc.add_paragraph(f"Nilai akhir estimasi: {skor['nilai_akhir']} dari {skor.get('nilai_akhir_maks')}")
    pred = skor.get("prediksi_peringkat") or {}
    doc.add_paragraph(f"Prediksi peringkat: {pred.get('nama', '-')}")
    doc.add_paragraph().add_run(CATATAN_ESTIMASI).italic = True

    for k in cuplikan["kriteria"]:
        bobot = f" (bobot {k['bobot']})" if infokom and k.get("bobot") is not None else ""
        doc.add_heading(f"Kriteria {k['kode']} — {k['nama']}{bobot}", 1)
        for b in k["butir"]:
            doc.add_heading(f"{b['kode']}. {b['elemen']}", 2)
            tabel = doc.add_table(rows=1, cols=4 if infokom else 3)
            tabel.style = "Light Grid Accent 1"
            judul = ["Skor estimasi", "Status", "Bobot", "Jenis / Tahap"] if infokom else ["Skor estimasi", "Status", "Syarat perlu"]
            for i, t in enumerate(judul):
                tabel.rows[0].cells[i].text = t
            sel = tabel.add_row().cells
            sel[0].text = str(b["skor_estimasi"] if b["skor_estimasi"] is not None else "-")
            sel[1].text = b["status_kelengkapan"]
            if infokom:
                sel[2].text = str(b["bobot"] if b["bobot"] is not None else "-")
                sel[3].text = " / ".join(x for x in [b.get("jenis_ipo"), b.get("tahap_ppepp")] if x) or "-"
            else:
                sel[2].text = "Ya" if any(a["is_syarat_perlu"] for a in b.get("ambang", [])) else "Tidak"
            for n in b["narasi"]:
                judul_bagian = next((s["judul"] for s in b["skema_narasi"] if s["bagian"] == n["bagian"]), n["bagian"])
                doc.add_paragraph().add_run(judul_bagian).bold = True
                doc.add_paragraph(n["isi"] or "(belum diisi)")
            if not b["narasi"]:
                doc.add_paragraph().add_run("(narasi belum diisi)").italic = True
            if b["data_terpetakan"]:
                doc.add_paragraph().add_run("Data pendukung (terpetakan otomatis)").bold = True
                for d in b["data_terpetakan"]:
                    doc.add_paragraph(teks_data(d), style="List Bullet")
            if b["bukti"]:
                doc.add_paragraph().add_run("Bukti").bold = True
                for bk in b["bukti"]:
                    doc.add_paragraph(teks_bukti(bk), style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
